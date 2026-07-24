"""WS2-3 selective export — build_export(only=) lọc theo check_code.

Không chọn = xuất toàn bộ; chọn = chỉ finding của mã đã chọn; sheet Tổng quan ghi
mã đã chọn (kỷ luật không-hộp-đen, ADR #18 Revision — WS2).
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.models import Finding
from app.pipeline.export import build_export


def _seed(session, company_id, code, subject):
    session.add(Finding(
        company_id=company_id, period_year=2024, check_code=code, severity="critical",
        subject_type="material_code", subject_key=subject, title=f"{code} {subject}",
    ))


def _finding_codes(xlsx: bytes) -> list[str]:
    wb = load_workbook(BytesIO(xlsx))
    ws = wb["Phát hiện"]
    return [row[0].value for row in ws.iter_rows(min_row=2) if row[0].value]


def _overview_text(xlsx: bytes) -> str:
    wb = load_workbook(BytesIO(xlsx))
    ws = wb["Tổng quan"]
    return "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )


def test_export_all_when_no_selection(session, company):
    _seed(session, company.id, "C1.1", "MAT_A")
    _seed(session, company.id, "C4.3", "MAT_B")
    session.commit()

    codes = _finding_codes(build_export(session, company, 2024))
    assert set(codes) == {"C1.1", "C4.3"}
    assert "Toàn bộ kiểm tra" in _overview_text(build_export(session, company, 2024))


def test_export_only_selected_codes(session, company):
    _seed(session, company.id, "C1.1", "MAT_A")
    _seed(session, company.id, "C4.3", "MAT_B")
    _seed(session, company.id, "C2.3", "MAT_C")
    session.commit()

    xlsx = build_export(session, company, 2024, only={"C1.1", "C4.3"})
    assert set(_finding_codes(xlsx)) == {"C1.1", "C4.3"}
    text = _overview_text(xlsx)
    assert "Các test đã chọn" in text
    assert "C1.1" in text and "C4.3" in text


def test_export_combo_code_is_selectable(session, company):
    _seed(session, company.id, "C1.1", "MAT_A")
    _seed(session, company.id, "COMBO_FORGED_NORM", "MAT_A")
    session.commit()

    xlsx = build_export(session, company, 2024, only={"COMBO_FORGED_NORM"})
    assert set(_finding_codes(xlsx)) == {"COMBO_FORGED_NORM"}


def test_export_route_binds_repeated_check_param():
    """FastAPI bind ?check=C1.1&check=C4.3 → list; tên file mang phạm vi đã chọn."""
    from fastapi.testclient import TestClient
    from sqlalchemy.pool import StaticPool

    import app.database as dbmod
    from app.auth_users import seed_default_admin
    from app.database import Base, SessionLocal, engine
    from app.main import app
    from app.models import Company

    new_engine = dbmod.create_engine(
        "sqlite://", connect_args={"check_same_thread": False},
        poolclass=StaticPool, future=True,
    )
    new_session = dbmod.sessionmaker(bind=new_engine, autoflush=False, autocommit=False, future=True)
    dbmod.engine = new_engine
    dbmod.SessionLocal = new_session
    Base.metadata.create_all(new_engine)
    try:
        with new_session() as db:
            seed_default_admin(db, "admin", "admin")
            c = Company(code="DN_EXP", name="Test", tax_id="1")
            db.add(c)
            db.flush()
            _seed(db, c.id, "C1.1", "MAT_A")
            _seed(db, c.id, "C4.3", "MAT_B")
            db.commit()
        client = TestClient(app)
        client.post("/login", data={"user": "admin", "password": "admin"}, follow_redirects=False)
        r = client.get("/companies/DN_EXP/export?year=2024&check=C1.1&check=C4.3")
        assert r.status_code == 200
        assert "C11-C43" in r.headers["content-disposition"]
        assert _finding_codes(r.content) == ["C1.1", "C4.3"]
    finally:
        new_engine.dispose()
        dbmod.engine = engine
        dbmod.SessionLocal = SessionLocal
