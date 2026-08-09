"""Trang tính (sheet) được đọc: ghi lại, hiện ra, và cho cán bộ chọn lại.

Adapter tự chấm điểm chọn trang, nhưng trước đây không ghi lại đã chọn trang nào
và màn xác nhận cột luôn vẽ trang ĐẦU workbook. File BCCT của 006 có `Tổng hợp`
đứng trước `Chi tiết` đang được nạp, nên cán bộ xác nhận chỉ số cột trên một trang
khác hẳn trang được đọc. Test giữ ba điều: registry ghi trang đã đọc, lưới xem
trước là trang đó, và ô chọn trang ghim được lựa chọn của cán bộ cho lượt nạp sau.
"""

from __future__ import annotations

import io
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy.pool import StaticPool

import app.database as dbmod
from app.database import Base, SessionLocal, engine
from app.main import app
from app.models import Company, DataFile, DataFileStatus, NvlBalance
from app.pipeline.file_page import file_page_url
from app.settings import settings
from tests.helpers import drain_jobs, upload_and_ingest

_M15_HEADER = [
    "STT", "Mã NVL", "Tên NVL", "Đơn vị tính", "Tồn đầu kỳ", "Nhập trong kỳ",
    "Tái xuất", "Chuyển mục đích sử dụng", "Xuất sản xuất", "Xuất khác", "Tồn cuối kỳ",
]


def _sheet_rows(ws, codes: list[str]) -> None:
    for _ in range(8):
        ws.append([None] * len(_M15_HEADER))
    ws.append(_M15_HEADER)
    for i, code in enumerate(codes):
        ws.append([i + 1, code, "Tên", "KG", 10, 100, 0, 0, 80, 0, 30])


def _two_sheet_m15() -> bytes:
    """Workbook có HAI trang cùng bố cục Mẫu 15, mã hàng khác nhau.

    Trang `BCQT_NVL` (tên khớp mẫu) đứng SAU trang phụ, nên "trang đầu workbook"
    và "trang được đọc" là hai trang khác nhau — đúng tình huống của file 006.
    """
    wb = Workbook()
    ws_extra = wb.active
    ws_extra.title = "Phụ lục"
    ws_extra.append(["Trang phụ, không phải biểu Mẫu 15"])
    _sheet_rows(wb.create_sheet("BCQT_NVL"), ["MAT0", "MAT1", "MAT2"])
    _sheet_rows(wb.create_sheet("BCQT_NVL_KY_SAU"), ["ALT0"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _setup(tmp_path: Path):
    import app.pipeline.ingest as ingmod
    from app.auth_users import seed_default_admin

    new_engine = dbmod.create_engine(
        "sqlite://", connect_args={"check_same_thread": False},
        poolclass=StaticPool, future=True,
    )
    new_session = dbmod.sessionmaker(
        bind=new_engine, autoflush=False, autocommit=False, future=True
    )
    dbmod.engine = new_engine
    dbmod.SessionLocal = new_session
    ingmod.SessionLocal = new_session
    Base.metadata.create_all(new_engine)
    with new_session() as db:
        seed_default_admin(db, "admin", "admin")
        db.add(Company(code="DN_SHEET", name="Sheet", tax_id="1"))
        db.commit()
    prev_root = settings.raw_data_path
    settings.raw_data_path = str(tmp_path)
    return new_engine, prev_root


def _teardown(new_engine, prev_root):
    import app.pipeline.ingest as ingmod
    settings.raw_data_path = prev_root
    new_engine.dispose()
    dbmod.engine = engine
    dbmod.SessionLocal = SessionLocal
    ingmod.SessionLocal = SessionLocal


def _login(client):
    client.post("/login", data={"user": "admin", "password": "admin"}, follow_redirects=False)


def _upload(client):
    r = upload_and_ingest(client, "DN_SHEET", "Mau15_NVL.xlsx", _two_sheet_m15())
    assert r.status_code == 303
    drain_jobs()


def _grid_mount(text: str) -> dict[str, str]:
    """Thuộc tính `data-*` của điểm neo lưới — hợp đồng giữa trang và lưới cuộn."""
    import html as _html
    import re

    tag = re.search(r"<div[^>]*id=\"cell-grid\"[^>]*>", text)
    assert tag, "trang không có điểm neo lưới"
    return {
        m.group(1): _html.unescape(m.group(3))
        for m in re.finditer(r"data-([a-z-]+)=([\"'])(.*?)\2", tag.group(0))
    }


def _m15_row():
    with dbmod.SessionLocal() as db:
        c = db.query(Company).filter_by(code="DN_SHEET").first()
        row = db.query(DataFile).filter_by(company_id=c.id, slot="m15").first()
        return row.id, row.parse_detail_obj, row.sheet_override


def _codes() -> set[str]:
    with dbmod.SessionLocal() as db:
        c = db.query(Company).filter_by(code="DN_SHEET").first()
        return {
            b.material_code
            for b in db.query(NvlBalance).filter_by(company_id=c.id, period_year=2024)
        }


def test_registry_records_the_sheet_that_was_read(tmp_path):
    new_engine, prev_root = _setup(tmp_path)
    try:
        client = TestClient(app)
        _login(client)
        _upload(client)

        _fid, detail, override = _m15_row()
        assert detail["sheet"] == "BCQT_NVL"   # KHÔNG phải "Phụ lục" (trang đầu)
        assert override is None                # tự nhận diện, chưa ai ghim
        assert _codes() == {"MAT0", "MAT1", "MAT2"}
    finally:
        _teardown(new_engine, prev_root)


def test_review_screen_previews_the_parsed_sheet_and_offers_the_others(tmp_path):
    new_engine, prev_root = _setup(tmp_path)
    try:
        client = TestClient(app)
        _login(client)
        _upload(client)
        fid, _detail, _ = _m15_row()

        html = client.get(file_page_url("DN_SHEET", fid)).text
        assert "Cài đặt đọc file" in html
        # Lưới mở ở trang ĐƯỢC ĐỌC, và mọi trang đều xem lại được (#124: bộ chọn đổi
        # thứ đang xem, ghim là nút riêng).
        attrs = _grid_mount(html)
        assert attrs["parsed-sheet"] == "BCQT_NVL"
        assert attrs["sheet"] == "1"           # "Phụ lục" đứng trước trong workbook
        assert 'data-sheet-name="Phụ lục"' in html
        assert 'data-sheet-name="BCQT_NVL_KY_SAU"' in html
        # Ô của trang được đọc đi qua điểm cuối cửa sổ, không dựng sẵn trong HTML.
        cells = client.get(
            f"/companies/DN_SHEET/documents/file/{fid}/cells?sheet=1"
        ).json()
        assert cells["sheet_name"] == "BCQT_NVL"
        assert any("MAT0" in [str(v) for v in row] for row in cells["rows"])
        assert "Trang phụ, không phải biểu Mẫu 15" not in html
    finally:
        _teardown(new_engine, prev_root)


def test_officer_can_pin_another_sheet_and_reingest_reads_it(tmp_path):
    new_engine, prev_root = _setup(tmp_path)
    try:
        client = TestClient(app)
        _login(client)
        _upload(client)
        fid, detail, _ = _m15_row()

        data = {f"col_{f}": str(i) for f, i in detail["column_map"].items()}
        data["sheet"] = "BCQT_NVL_KY_SAU"
        r = client.post(f"/companies/DN_SHEET/documents/file/{fid}/review",
                        data=data, follow_redirects=False)
        assert r.status_code == 303
        drain_jobs()

        _fid, detail_after, override = _m15_row()
        assert override == "BCQT_NVL_KY_SAU"          # ghim lại cho lượt nạp sau
        assert detail_after["sheet"] == "BCQT_NVL_KY_SAU"
        assert _codes() == {"ALT0"}                    # dòng đọc từ trang đã ghim
    finally:
        _teardown(new_engine, prev_root)


def _unreadable_m15() -> bytes:
    """Workbook mà `select_sheet` KHÔNG nhận ra trang nào đúng biểu.

    Dựng lại tình huống file cán bộ tự gộp: chèn một cột đầu (nhãn nguồn) và bỏ khối
    tiêu đề → mọi cột lệch một ô, mọi trang chấm 0 điểm, file bị từ chối khi nạp.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Bìa"
    ws.append(["Trang bìa"])
    ws2 = wb.create_sheet("Dữ liệu")
    ws2.append(["Nguồn", *_M15_HEADER])
    for i in range(3):
        ws2.append(["F1", i + 1, f"MAT{i}", "Tên", "KG", 10, 100, 0, 0, 80, 0, 30])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_unreadable_file_still_offers_the_sheet_picker(tmp_path):
    """File đọc hỏng vẫn phải ghim được trang tính — nếu không thì nó bế tắc.

    Không có bố cục cột để xác nhận (parse thất bại nên chưa có `column_map`), mà cũng
    không có cách chỉ cho hệ thống trang cần đọc → cán bộ chỉ còn cách sửa file nguồn.
    """
    new_engine, prev_root = _setup(tmp_path)
    try:
        client = TestClient(app)
        _login(client)
        upload_and_ingest(client, "DN_SHEET", "Mau15_NVL.xlsx", _unreadable_m15())
        drain_jobs()

        with dbmod.SessionLocal() as db:
            c = db.query(Company).filter_by(code="DN_SHEET").first()
            row = db.query(DataFile).filter_by(company_id=c.id, slot="m15").first()
            assert row.parse_status == DataFileStatus.ERROR   # đúng là file hỏng
            fid = row.id

        # Trang tài liệu phải có lối vào, và màn review phải có ô chọn trang.
        docs = client.get("/companies/DN_SHEET/documents").text
        assert f'href="{file_page_url("DN_SHEET", fid)}"' in docs
        html = client.get(file_page_url("DN_SHEET", fid)).text
        assert 'id="sheet-pin"' in html
        assert "chưa xác định trang nào chứa biểu" in html
        assert 'data-sheet-name="Dữ liệu"' in html

        # Ghim trang → nạp lại → file đọc được. Địa chỉ cũ vẫn nhận ô trang tính.
        r = client.post(f"/companies/DN_SHEET/documents/file/{fid}/review",
                        data={"sheet": "Dữ liệu"}, follow_redirects=False)
        assert r.status_code == 303
        assert r.headers["location"].endswith("/documents#ky-2024")
        drain_jobs()

        with dbmod.SessionLocal() as db:
            c = db.query(Company).filter_by(code="DN_SHEET").first()
            row = db.query(DataFile).filter_by(company_id=c.id, slot="m15").first()
            assert row.sheet_override == "Dữ liệu"
            assert row.parse_status == DataFileStatus.OK
            assert row.row_count == 3
    finally:
        _teardown(new_engine, prev_root)


def test_confirming_without_touching_the_sheet_keeps_auto_detection(tmp_path):
    """Bấm xác nhận mà không đụng ô trang tính KHÔNG được lặng lẽ ghim trang."""
    new_engine, prev_root = _setup(tmp_path)
    try:
        client = TestClient(app)
        _login(client)
        _upload(client)
        fid, detail, _ = _m15_row()

        data = {f"col_{f}": str(i) for f, i in detail["column_map"].items()}
        data["sheet"] = ""                             # ô để trống = tự nhận diện
        client.post(f"/companies/DN_SHEET/documents/file/{fid}/review",
                    data=data, follow_redirects=False)
        drain_jobs()

        _fid, detail_after, override = _m15_row()
        assert override is None
        assert detail_after["sheet"] == "BCQT_NVL"
        assert _codes() == {"MAT0", "MAT1", "MAT2"}
    finally:
        _teardown(new_engine, prev_root)
