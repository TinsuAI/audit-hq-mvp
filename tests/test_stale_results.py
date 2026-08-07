"""T10 (#90) — xoá/thay file dời phiên bản dữ liệu; dấu hiệu cũ ở cả ba nơi.

Hai TRỤC khác nhau, không trộn:

- **Trục bộ file** (`readiness.stale`): dòng Tầng 1 đã nạp không còn khớp bản ghi
  file của kỳ → việc cần làm là NẠP LẠI. Gỡ bằng lượt nạp.
- **Trục kết quả** (`results_stale`): `check_runs.data_version` còn ở sau
  `CompanyPeriod.data_version` → phát hiện và điểm rủi ro tính trên bộ dữ liệu cũ
  → việc cần làm là CHẠY LẠI KIỂM TRA. Cùng một phép so mà tổng quan AI đang dùng.

Không trục nào được chạm vào phép đếm "đủ dữ liệu cho N/M kiểm tra" (#85), và
không trục nào được TỰ chạy lại kiểm tra: `run_checks` xoá rồi dựng lại `Finding`,
đưa `status`/`notes` cán bộ đã đánh về "mới" (ADR #24).

Khẳng định ở mức DỮ LIỆU — trạng thái DB, ngữ cảnh màn hình, thuộc tính `data-*`
máy đọc được — chứ không dò chuỗi tiếng Việt trong HTML: các template ở đây đã có
test dò chữ, mà test dò chữ vẫn xanh sau khi màn hình đổi hết ý nghĩa.
"""

from __future__ import annotations

import io
from datetime import datetime

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select

from app.main import app
from app.models import (
    CheckRun,
    Company,
    CompanyPeriod,
    CompanyYearScore,
    DataFile,
    Finding,
    Job,
    NvlBalance,
)
from app.models.data_file import DataFileStatus
from app.pipeline.data_screen import ACTION_INGEST, build_data_screen
from app.pipeline.period import bump_data_version, current_data_version
from app.pipeline.staleness import results_stale, stale_result_years
from tests.conftest import AppDb, add_nvl
from tests.helpers import XLSX_MIME
from tests.test_readiness import add_file

_CODE = "DN_T10"


def _xlsx_bytes(marker: str) -> bytes:
    wb = Workbook()
    wb.active["A1"] = marker
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _client(app_db: AppDb, code: str = _CODE) -> TestClient:
    with app_db.SessionLocal() as db:
        db.add(Company(code=code, name="Kết quả cũ", tax_id="1"))
        db.commit()
    client = TestClient(app)
    client.post(
        "/login", data={"user": "admin", "password": "admin"}, follow_redirects=False
    )
    return client


def _company_id(app_db: AppDb, code: str = _CODE) -> int:
    with app_db.SessionLocal() as db:
        return db.scalar(select(Company.id).where(Company.code == code))


def add_run(session, company_id: int, *, year: int = 2024, code: str = "C1.1",
            data_version: int = 0) -> CheckRun:
    row = CheckRun(
        company_id=company_id, period_year=year, check_code=code,
        ran_at=datetime(2026, 8, 7, 10, 0), finding_count=0, status="ok",
        data_version=data_version,
    )
    session.add(row)
    return row


# --- Dời phiên bản dữ liệu -------------------------------------------------------


def test_bumping_creates_the_period_row_when_the_company_has_none(session, company):
    """Kỳ chưa có dòng `company_periods` vẫn phải dời được — phiên bản 0 → 1."""
    assert current_data_version(session, company.id, 2024) == 0

    assert bump_data_version(session, company.id, 2024) == 1
    session.commit()

    assert current_data_version(session, company.id, 2024) == 1


def test_bumping_moves_an_existing_version_by_one(session, company):
    session.add(CompanyPeriod(company_id=company.id, period_year=2024, data_version=7))
    session.commit()

    assert bump_data_version(session, company.id, 2024) == 8
    session.commit()

    assert current_data_version(session, company.id, 2024) == 8


def test_bumping_one_period_leaves_the_other_periods_alone(session, company):
    session.add(CompanyPeriod(company_id=company.id, period_year=2025, data_version=3))
    session.commit()

    bump_data_version(session, company.id, 2024)
    session.commit()

    assert current_data_version(session, company.id, 2025) == 3


# --- Trục kết quả: so phiên bản dữ liệu với lần chạy ------------------------------


def test_results_are_not_stale_before_any_check_has_run(session, company):
    """Dòng `check_runs` VẮNG = "chưa rõ", KHÔNG phải cũ (WS3)."""
    session.add(CompanyPeriod(company_id=company.id, period_year=2024, data_version=4))
    session.commit()

    assert results_stale(session, company.id, 2024) is False


def test_results_are_current_when_the_run_read_the_current_version(session, company):
    session.add(CompanyPeriod(company_id=company.id, period_year=2024, data_version=2))
    add_run(session, company.id, data_version=2)
    session.commit()

    assert results_stale(session, company.id, 2024) is False


def test_results_are_stale_once_the_data_version_moves_past_the_run(session, company):
    session.add(CompanyPeriod(company_id=company.id, period_year=2024, data_version=2))
    add_run(session, company.id, data_version=1)
    session.commit()

    assert results_stale(session, company.id, 2024) is True


def test_a_partial_rerun_leaves_the_period_stale(session, company):
    """Chạy lại một mã không làm mới cả kỳ: điểm gộp phát hiện của hai phiên bản."""
    session.add(CompanyPeriod(company_id=company.id, period_year=2024, data_version=3))
    add_run(session, company.id, code="C1.1", data_version=3)
    add_run(session, company.id, code="C1.2", data_version=2)
    session.commit()

    assert results_stale(session, company.id, 2024) is True


def test_staleness_is_per_period_not_per_company(session, company):
    session.add_all([
        CompanyPeriod(company_id=company.id, period_year=2024, data_version=2),
        CompanyPeriod(company_id=company.id, period_year=2025, data_version=1),
    ])
    add_run(session, company.id, year=2024, data_version=1)
    add_run(session, company.id, year=2025, data_version=1)
    session.commit()

    assert results_stale(session, company.id, 2024) is True
    assert results_stale(session, company.id, 2025) is False


def test_stale_result_years_lists_every_company_with_results_behind(session, company):
    other = Company(code="DN_KHAC", name="DN khác", tax_id="2")
    session.add(other)
    session.flush()
    session.add_all([
        CompanyPeriod(company_id=company.id, period_year=2024, data_version=2),
        CompanyPeriod(company_id=company.id, period_year=2025, data_version=5),
        CompanyPeriod(company_id=other.id, period_year=2024, data_version=1),
    ])
    add_run(session, company.id, year=2024, data_version=1)
    add_run(session, company.id, year=2025, data_version=5)
    add_run(session, other.id, year=2024, data_version=1)
    session.commit()

    years = stale_result_years(session)
    assert years[company.id] == (2024,)
    assert other.id not in years


# --- Xoá file: dời phiên bản, KHÔNG đụng dòng, KHÔNG chạy lại ---------------------


def _upload(client: TestClient, slot: str, name: str, code: str = _CODE, year: int = 2024):
    return client.post(
        f"/companies/{code}/documents/upload",
        data={"year": str(year), "slot": slot},
        files=[("file", (name, _xlsx_bytes(name), XLSX_MIME))],
        follow_redirects=False,
    )


def test_deleting_a_file_moves_the_period_data_version(app_db: AppDb):
    client = _client(app_db)
    assert _upload(client, "m15", "Mau15.xlsx").status_code == 303

    cid = _company_id(app_db)
    with app_db.SessionLocal() as db:
        before = current_data_version(db, cid, 2024)
        file_id = db.scalar(select(DataFile.id).where(DataFile.company_id == cid))

    r = client.post(
        f"/companies/{_CODE}/documents/file/{file_id}/delete", follow_redirects=False
    )
    assert r.status_code == 303

    with app_db.SessionLocal() as db:
        assert current_data_version(db, cid, 2024) == before + 1
        assert db.scalar(select(DataFile.id).where(DataFile.company_id == cid)) is None


def test_deleting_a_file_keeps_the_loaded_rows_and_the_officer_review(app_db: AppDb):
    """Chạy lại kiểm tra xoá rồi dựng lại `Finding` → mất `status`/`notes` cán bộ đã
    đánh. Xoá file KHÔNG được kéo theo lượt chạy nào."""
    client = _client(app_db)
    assert _upload(client, "m15", "Mau15.xlsx").status_code == 303
    cid = _company_id(app_db)

    with app_db.SessionLocal() as db:
        add_nvl(db, cid, material_code="A", imported=10, closing=10, year=2024)
        db.add(Finding(
            company_id=cid, period_year=2024, check_code="C1.1", severity="warning",
            title="Chênh lệch", status="reviewed", notes="Cán bộ đã đối chiếu",
        ))
        add_run(db, cid, data_version=1)
        db.commit()
        file_id = db.scalar(select(DataFile.id).where(DataFile.company_id == cid))
        jobs_before = db.scalar(select(Job.id).where(Job.company_id == cid).limit(1))

    client.post(
        f"/companies/{_CODE}/documents/file/{file_id}/delete", follow_redirects=False
    )

    with app_db.SessionLocal() as db:
        assert db.scalar(
            select(NvlBalance.id).where(NvlBalance.company_id == cid)
        ) is not None
        finding = db.scalar(select(Finding).where(Finding.company_id == cid))
        assert (finding.status, finding.notes) == ("reviewed", "Cán bộ đã đối chiếu")
        run = db.scalar(select(CheckRun).where(CheckRun.company_id == cid))
        assert run.data_version == 1
        assert db.scalar(
            select(Job.id).where(Job.company_id == cid).limit(1)
        ) == jobs_before


def test_deleting_a_file_does_not_change_how_many_checks_have_enough_data(app_db: AppDb):
    """Trục riêng: sau khi xoá, "đủ dữ liệu" vẫn đúng — kiểm tra vẫn chạy được (#85)."""
    client = _client(app_db)
    assert _upload(client, "m15", "Mau15.xlsx").status_code == 303
    cid = _company_id(app_db)

    with app_db.SessionLocal() as db:
        add_nvl(db, cid, material_code="A", imported=10, closing=10, year=2024)
        db.commit()
        company = db.get(Company, cid)
        before = next(
            p for p in build_data_screen(db, company).periods if p.year == 2024
        )
        counts_before = (before.sufficient_count, before.total_count)
        file_id = db.scalar(select(DataFile.id).where(DataFile.company_id == cid))

    client.post(
        f"/companies/{_CODE}/documents/file/{file_id}/delete", follow_redirects=False
    )

    with app_db.SessionLocal() as db:
        company = db.get(Company, cid)
        after = next(p for p in build_data_screen(db, company).periods if p.year == 2024)
        assert (after.sufficient_count, after.total_count) == counts_before
        assert after.stale is True


def test_replacing_a_file_moves_the_period_data_version(app_db: AppDb):
    """Ô một-file: file mới thay file cũ → bộ file đã đổi, dòng đã nạp là của bộ trước."""
    client = _client(app_db)
    assert _upload(client, "m15", "Mau15_cu.xlsx").status_code == 303
    cid = _company_id(app_db)
    with app_db.SessionLocal() as db:
        before = current_data_version(db, cid, 2024)

    assert _upload(client, "m15", "Mau15_moi.xlsx").status_code == 303

    with app_db.SessionLocal() as db:
        assert current_data_version(db, cid, 2024) == before + 1


def test_replacing_a_bcct_file_by_name_moves_the_period_data_version(app_db: AppDb):
    client = _client(app_db)
    assert _upload(client, "bcct", "BCCT_F1.xlsx").status_code == 303
    cid = _company_id(app_db)
    with app_db.SessionLocal() as db:
        before = current_data_version(db, cid, 2024)

    assert _upload(client, "bcct", "BCCT_F1.xlsx").status_code == 303

    with app_db.SessionLocal() as db:
        assert current_data_version(db, cid, 2024) == before + 1


def test_adding_a_brand_new_file_does_not_move_the_data_version(app_db: AppDb):
    """Thêm file KHÔNG làm dòng đã nạp sai đi — trục bộ file của `readiness` lo việc
    đó. Dời phiên bản ở đây là báo phát hiện cũ trong khi chưa gì đổi."""
    client = _client(app_db)
    assert _upload(client, "bcct", "BCCT_F1.xlsx").status_code == 303
    cid = _company_id(app_db)
    with app_db.SessionLocal() as db:
        before = current_data_version(db, cid, 2024)

    assert _upload(client, "bcct", "BCCT_F3.xlsx").status_code == 303

    with app_db.SessionLocal() as db:
        assert current_data_version(db, cid, 2024) == before


# --- Nơi 1: dòng kỳ --------------------------------------------------------------


def test_the_period_row_says_the_file_set_changed_in_one_line(session, company):
    """Ba loại tài liệu cùng cũ là MỘT việc (nạp lại kỳ), không phải ba việc."""
    add_nvl(session, company.id, material_code="A", imported=1, closing=1, year=2024)
    from tests.conftest import add_norm, add_sp
    add_sp(session, company.id, product_code="P", intake=1, closing=1, year=2024)
    add_norm(
        session, company.id, product_code="P", material_code="A", norm_qty=1, year=2024,
    )
    session.commit()

    row = next(p for p in build_data_screen(session, company).periods if p.year == 2024)
    stale_items = [
        item for group in row.groups for item in group.items if item.slots
    ]
    assert len(stale_items) == 1
    assert stale_items[0].slots == ("m15", "m15a", "m16")
    assert stale_items[0].action.kind == ACTION_INGEST
    assert row.stale is True


def test_one_stale_document_type_still_produces_exactly_one_line(session, company):
    add_nvl(session, company.id, material_code="A", imported=1, closing=1, year=2024)
    session.commit()

    row = next(p for p in build_data_screen(session, company).periods if p.year == 2024)
    stale_items = [item for group in row.groups for item in group.items if item.slots]
    assert len(stale_items) == 1
    assert stale_items[0].slots == ("m15",)


def test_the_period_row_marks_results_computed_on_an_older_data_version(session, company):
    add_file(session, company.id, slot="m15", status=DataFileStatus.OK, year=2024)
    add_nvl(session, company.id, material_code="A", imported=1, closing=1, year=2024)
    session.add(CompanyPeriod(company_id=company.id, period_year=2024, data_version=2))
    add_run(session, company.id, year=2024, data_version=1)
    session.add(CompanyYearScore(
        company_id=company.id, period_year=2024, score=129, tier="Cao", breakdown={},
    ))
    session.commit()

    row = next(p for p in build_data_screen(session, company).periods if p.year == 2024)
    # Bộ file khớp lại rồi (đã nạp lại) nhưng kết quả vẫn của lần chạy trước.
    assert row.stale is False
    assert row.results_stale is True
    assert row.score == 129  # KHÔNG giấu số


def test_the_period_row_has_no_result_mark_when_the_run_is_current(session, company):
    add_file(session, company.id, slot="m15", status=DataFileStatus.OK, year=2024)
    add_nvl(session, company.id, material_code="A", imported=1, closing=1, year=2024)
    session.add(CompanyPeriod(company_id=company.id, period_year=2024, data_version=2))
    add_run(session, company.id, year=2024, data_version=2)
    session.commit()

    row = next(p for p in build_data_screen(session, company).periods if p.year == 2024)
    assert row.results_stale is False


# --- Nơi 2: màn phát hiện ---------------------------------------------------------


def _seed_stale_results(app_db: AppDb, code: str = _CODE, *, score: int = 129,
                        run_version: int = 1, data_version: int = 2) -> int:
    cid = _company_id(app_db, code)
    with app_db.SessionLocal() as db:
        add_nvl(db, cid, material_code="A", imported=10, closing=10, year=2024)
        db.add(CompanyPeriod(
            company_id=cid, period_year=2024, data_version=data_version,
        ))
        db.add(CompanyYearScore(
            company_id=cid, period_year=2024, score=score, tier="Cao", breakdown={},
        ))
        db.add(Finding(
            company_id=cid, period_year=2024, check_code="C1.1", severity="warning",
            title="Chênh lệch",
        ))
        add_run(db, cid, year=2024, data_version=run_version)
        db.commit()
    return cid


def test_the_findings_screen_marks_results_from_an_older_data_version(app_db: AppDb):
    client = _client(app_db)
    _seed_stale_results(app_db)

    html = client.get(f"/companies/{_CODE}?year=2024").text
    assert 'data-results-stale="1"' in html
    # KHÔNG giấu số: điểm vẫn hiện nguyên.
    assert "129" in html


def test_the_findings_screen_has_no_mark_when_the_run_is_current(app_db: AppDb):
    client = _client(app_db)
    _seed_stale_results(app_db, run_version=2, data_version=2)

    html = client.get(f"/companies/{_CODE}?year=2024").text
    assert 'data-results-stale="1"' not in html


# --- Nơi 3: cột điểm ở bảng danh sách DN -----------------------------------------


def _row_attrs(html: str, code: str) -> str:
    """Thẻ `<tr>` của một DN — neo vào đường dẫn chi tiết trong chính dòng đó."""
    import re

    for block in re.findall(r"<tr class=\"dn-row.*?</tr>", html, flags=re.S):
        if f"/companies/{code}" in block:
            return block
    raise AssertionError(f"Không tìm thấy dòng của {code}")


def test_a_company_with_stale_results_stays_in_the_ranking_with_a_mark(app_db: AppDb):
    """Rơi khỏi xếp hạng vì có người tải file lên là cùng dạng sai lầm với việc bỏ
    luật khỏi thang điểm (#65) — nhãn, không phải gỡ khỏi bảng."""
    client = _client(app_db)
    with app_db.SessionLocal() as db:
        db.add(Company(code="DN_MOI", name="DN mới", tax_id="2"))
        db.commit()
    _seed_stale_results(app_db, _CODE, score=129, run_version=1, data_version=2)
    _seed_stale_results(app_db, "DN_MOI", score=54, run_version=2, data_version=2)

    html = client.get("/companies").text
    stale_row = _row_attrs(html, _CODE)
    fresh_row = _row_attrs(html, "DN_MOI")

    assert 'data-stale="1"' in stale_row
    assert 'data-stale="0"' in fresh_row
    # Điểm vẫn nằm trong dòng và vẫn sắp xếp được bằng chính giá trị đó.
    assert 'data-score="129"' in stale_row
