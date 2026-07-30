"""NIENDO-1 (#50, ADR #23 T2) — niên độ là thuộc tính DN.

`companies.fiscal_start_month` ∈ {1, 4, 7, 10} (điểm a khoản 1 Điều 12 Luật Kế toán
88/2015: niên độ khác dương lịch phải 12 tháng tròn tính từ đầu quý). Nhãn năm = năm
BẮT ĐẦU kỳ; pháp luật định danh kỳ CHỈ bằng khoảng ngày nên nhãn là khoá nội bộ và
mọi màn hiện nhãn kỳ ≠ dương lịch phải in kèm khoảng ngày.

Chuỗi suy cửa sổ: is_manual > tiêu đề file (đủ 2 ngày) > default niên độ DN > dương lịch.
"""

from __future__ import annotations

from datetime import date

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.adapters._common import CompanyHeader
from app.main import app
from app.models import Company, CompanyPeriod
from app.pipeline.period import (
    FISCAL_START_MONTHS,
    default_bounds,
    fiscal_bounds,
    header_conflicts_with_fiscal_year,
    load_period_windows,
    resolve_period_bounds,
)
from tests.test_company_period_route import _login, _setup_db, _teardown

# --- Cửa sổ mặc định từ niên độ ----------------------------------------------


def test_calendar_default_is_unchanged():
    assert fiscal_bounds(2025, 1) == (date(2025, 1, 1), date(2025, 12, 31))


def test_april_fiscal_year_labels_by_start_year():
    assert fiscal_bounds(2025, 4) == (date(2025, 4, 1), date(2026, 3, 31))


def test_july_and_october_fiscal_years():
    assert fiscal_bounds(2025, 7) == (date(2025, 7, 1), date(2026, 6, 30))
    assert fiscal_bounds(2025, 10) == (date(2025, 10, 1), date(2026, 9, 30))


def test_only_quarter_starts_are_accepted():
    assert FISCAL_START_MONTHS == (1, 4, 7, 10)
    with pytest.raises(ValueError):
        fiscal_bounds(2025, 2)


# --- Chuỗi ưu tiên ------------------------------------------------------------


def test_default_bounds_uses_fiscal_start_month_when_header_incomplete():
    assert default_bounds(None, 2025, fiscal_start_month=4) == (
        date(2025, 4, 1), date(2026, 3, 31)
    )


def test_header_with_both_dates_beats_fiscal_default():
    h = CompanyHeader(period_from=date(2025, 1, 1), period_to=date(2025, 12, 31))
    assert default_bounds(h, 2025, fiscal_start_month=4) == (
        date(2025, 1, 1), date(2025, 12, 31)
    )


def test_calendar_fallback_when_company_has_no_fiscal_setting():
    assert default_bounds(None, 2025) == (date(2025, 1, 1), date(2025, 12, 31))


def test_manual_window_beats_everything(session, company):
    company.fiscal_start_month = 4
    session.add(CompanyPeriod(
        company_id=company.id, period_year=2025,
        period_from=date(2025, 7, 1), period_to=date(2026, 6, 30), is_manual=True,
    ))
    session.commit()
    h = CompanyHeader(period_from=date(2025, 1, 1), period_to=date(2025, 12, 31))
    assert resolve_period_bounds(session, company.id, 2025, h) == (
        date(2025, 7, 1), date(2026, 6, 30)
    )


def test_resolve_uses_company_fiscal_default_when_no_row_and_no_header(session, company):
    company.fiscal_start_month = 10
    session.commit()
    assert resolve_period_bounds(session, company.id, 2025) == (
        date(2025, 10, 1), date(2026, 9, 30)
    )
    row = session.scalar(
        select(CompanyPeriod).where(
            CompanyPeriod.company_id == company.id, CompanyPeriod.period_year == 2025
        )
    )
    assert row is not None and row.is_manual is False


def test_companies_without_setting_keep_calendar_behaviour(session, company):
    assert company.fiscal_start_month == 1
    assert resolve_period_bounds(session, company.id, 2025) == (
        date(2025, 1, 1), date(2025, 12, 31)
    )


# --- Cảnh báo tiêu đề lệch niên độ -------------------------------------------


def test_header_matching_fiscal_default_is_no_conflict():
    h = CompanyHeader(period_from=date(2025, 4, 1), period_to=date(2026, 3, 31))
    assert header_conflicts_with_fiscal_year(h, 2025, 4) is None


def test_header_conflicting_with_fiscal_default_is_reported():
    h = CompanyHeader(period_from=date(2025, 1, 1), period_to=date(2025, 12, 31))
    conflict = header_conflicts_with_fiscal_year(h, 2025, 4)
    assert conflict == ((date(2025, 1, 1), date(2025, 12, 31)),
                        (date(2025, 4, 1), date(2026, 3, 31)))


def test_no_conflict_reported_for_calendar_companies():
    h = CompanyHeader(period_from=date(2025, 1, 1), period_to=date(2025, 12, 31))
    assert header_conflicts_with_fiscal_year(h, 2025, 1) is None


def test_incomplete_header_cannot_conflict():
    assert header_conflicts_with_fiscal_year(CompanyHeader(period_from=date(2025, 1, 1)),
                                             2025, 4) is None
    assert header_conflicts_with_fiscal_year(None, 2025, 4) is None


# --- Nhãn kỳ kèm khoảng ngày ở mọi màn ---------------------------------------


def test_load_period_windows_covers_years_without_a_stored_row(session, company):
    company.fiscal_start_month = 4
    session.commit()
    windows = load_period_windows(session, company.id, years=[2024, 2025])
    assert windows == {
        2024: (date(2024, 4, 1), date(2025, 3, 31)),
        2025: (date(2025, 4, 1), date(2026, 3, 31)),
    }


def test_load_period_windows_stays_empty_for_calendar_companies(session, company):
    assert load_period_windows(session, company.id, years=[2024, 2025]) == {}


def test_stored_row_wins_over_company_default(session, company):
    company.fiscal_start_month = 4
    session.add(CompanyPeriod(
        company_id=company.id, period_year=2025,
        period_from=date(2025, 7, 1), period_to=date(2026, 6, 30), is_manual=True,
    ))
    session.commit()
    assert load_period_windows(session, company.id, years=[2025]) == {
        2025: (date(2025, 7, 1), date(2026, 6, 30))
    }


# --- UI cài đặt DN ------------------------------------------------------------


def test_edit_company_page_offers_the_four_quarter_starts():
    new_engine, _ = _setup_db()
    try:
        client = TestClient(app)
        _login(client)
        r = client.get("/companies/DN_077/edit")
        assert r.status_code == 200
        assert "Niên độ kế toán" in r.text
        for label in ("01/01", "01/04", "01/07", "01/10"):
            assert label in r.text
    finally:
        _teardown(new_engine)


def test_saving_fiscal_start_month_persists():
    new_engine, new_session = _setup_db()
    try:
        client = TestClient(app)
        _login(client)
        r = client.post(
            "/companies/DN_077/edit",
            data={"name": "Cơ khí Test", "tax_id": "111", "address": "",
                  "industry": "", "fiscal_start_month": "4"},
            follow_redirects=False,
        )
        assert r.status_code == 303
        with new_session() as db:
            c = db.scalar(select(Company).where(Company.code == "DN_077"))
            assert c.fiscal_start_month == 4
    finally:
        _teardown(new_engine)


def test_invalid_fiscal_start_month_is_rejected():
    new_engine, new_session = _setup_db()
    try:
        client = TestClient(app)
        _login(client)
        r = client.post(
            "/companies/DN_077/edit",
            data={"name": "Cơ khí Test", "tax_id": "111", "address": "",
                  "industry": "", "fiscal_start_month": "2"},
            follow_redirects=False,
        )
        assert r.status_code == 303
        assert "error" in r.headers["location"]
        with new_session() as db:
            c = db.scalar(select(Company).where(Company.code == "DN_077"))
            assert c.fiscal_start_month == 1
    finally:
        _teardown(new_engine)


def test_export_header_prints_the_date_range_for_fiscal_years():
    from io import BytesIO

    from openpyxl import load_workbook

    from app.pipeline.export import build_export

    new_engine, new_session = _setup_db()
    try:
        with new_session() as db:
            c = db.scalar(select(Company).where(Company.code == "DN_077"))
            c.fiscal_start_month = 4
            db.commit()
            payload = build_export(db, c, 2025)
        ws = load_workbook(BytesIO(payload))["Tổng quan"]
        cells = [ws.cell(row=r, column=2).value for r in range(1, 12)]
        assert any(v and "01/04/2025" in str(v) and "31/03/2026" in str(v) for v in cells)
    finally:
        _teardown(new_engine)


# --- Cảnh báo trên màn review -------------------------------------------------


def test_stored_window_matching_fiscal_default_is_no_conflict(session, company):
    from app.pipeline.period import period_window_conflict

    company.fiscal_start_month = 4
    session.add(CompanyPeriod(
        company_id=company.id, period_year=2025,
        period_from=date(2025, 4, 1), period_to=date(2026, 3, 31), is_manual=False,
    ))
    session.commit()
    assert period_window_conflict(session, company.id, 2025) is None


def test_stored_header_window_conflicting_with_fiscal_default_is_reported(session, company):
    from app.pipeline.period import period_window_conflict

    company.fiscal_start_month = 4
    session.add(CompanyPeriod(  # tiêu đề file nói dương lịch, DN khai niên độ tháng 4
        company_id=company.id, period_year=2025,
        period_from=date(2025, 1, 1), period_to=date(2025, 12, 31), is_manual=False,
    ))
    session.commit()
    assert period_window_conflict(session, company.id, 2025) == (
        (date(2025, 1, 1), date(2025, 12, 31)),
        (date(2025, 4, 1), date(2026, 3, 31)),
    )


def test_manual_window_is_a_deliberate_choice_not_a_conflict(session, company):
    from app.pipeline.period import period_window_conflict

    company.fiscal_start_month = 4
    session.add(CompanyPeriod(
        company_id=company.id, period_year=2025,
        period_from=date(2025, 1, 1), period_to=date(2025, 12, 31), is_manual=True,
    ))
    session.commit()
    assert period_window_conflict(session, company.id, 2025) is None


def test_calendar_company_never_conflicts(session, company):
    from app.pipeline.period import period_window_conflict

    session.add(CompanyPeriod(
        company_id=company.id, period_year=2025,
        period_from=date(2025, 4, 1), period_to=date(2026, 3, 31), is_manual=False,
    ))
    session.commit()
    assert period_window_conflict(session, company.id, 2025) is None


def test_review_screen_warns_when_file_window_conflicts_with_fiscal_year():
    from app.models import DataFile

    new_engine, new_session = _setup_db()
    try:
        with new_session() as db:
            c = db.scalar(select(Company).where(Company.code == "DN_077"))
            c.fiscal_start_month = 4
            db.add(CompanyPeriod(
                company_id=c.id, period_year=2025,
                period_from=date(2025, 1, 1), period_to=date(2025, 12, 31), is_manual=False,
            ))
            db.add(DataFile(
                company_id=c.id, period_year=2025, slot="m15",
                original_filename="NVL.xlsx", stored_path="DN_077/2025/BCQT/NVL.xlsx",
            ))
            db.commit()
            file_id = db.scalar(select(DataFile.id))
        client = TestClient(app)
        _login(client)
        r = client.get(f"/companies/DN_077/documents/file/{file_id}/review")
        assert r.status_code == 200
        assert "lệch với niên độ" in r.text
        assert "01/04/2025 – 31/03/2026" in r.text
    finally:
        _teardown(new_engine)
