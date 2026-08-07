"""Tests — registry data_files: sync filesystem + record parse result."""

from __future__ import annotations

from app.models import DataFile, DataFileStatus
from app.pipeline.data_files import (
    _classify_slot,
    files_by_year_slot,
    record_parse_result,
    sync_data_files,
)
from app.pipeline.ingest import IngestStats


def _touch(root, code, year, subdir, name, content=b"x"):
    d = root / code / str(year) / subdir
    d.mkdir(parents=True, exist_ok=True)
    p = d / name
    p.write_bytes(content)
    return p


class TestClassifySlot:
    def test_bcqt_nvl(self):
        assert _classify_slot("BCQT", "Mau15_NVL_2024.xlsx") == "m15"

    def test_bcqt_sp(self):
        assert _classify_slot("BCQT", "Mau15a_SP_2024.xlsx") == "m15a"

    def test_dinh_muc(self):
        assert _classify_slot("DINH_MUC", "BCDM_TT39.xls") == "m16"

    def test_hang_chi_tiet(self):
        assert _classify_slot("HANG_CHI_TIET", "BCCT_2024.xlsx") == "bcct"

    def test_unknown_bcqt(self):
        assert _classify_slot("BCQT", "old_tt38_file.xlsx") is None


class TestSync:
    def test_sync_creates_registry(self, session, company, tmp_path):
        _touch(tmp_path, company.code, 2024, "BCQT", "Mau15_NVL.xlsx")
        _touch(tmp_path, company.code, 2024, "BCQT", "Mau15a_SP.xlsx")
        _touch(tmp_path, company.code, 2024, "DINH_MUC", "BCDM_TT39.xls")
        _touch(tmp_path, company.code, 2023, "HANG_CHI_TIET", "BCCT.xlsx")

        sync_data_files(session, company, raw_root=tmp_path)

        rows = session.query(DataFile).filter_by(company_id=company.id).all()
        slots = sorted((r.period_year, r.slot) for r in rows)
        assert slots == [(2023, "bcct"), (2024, "m15"), (2024, "m15a"), (2024, "m16")]
        assert all(r.parse_status == DataFileStatus.PENDING for r in rows)
        assert all(r.size_bytes > 0 for r in rows)

    def test_sync_prunes_deleted(self, session, company, tmp_path):
        p = _touch(tmp_path, company.code, 2024, "BCQT", "Mau15_NVL.xlsx")
        sync_data_files(session, company, raw_root=tmp_path)
        assert session.query(DataFile).filter_by(company_id=company.id).count() == 1

        p.unlink()
        sync_data_files(session, company, raw_root=tmp_path)
        assert session.query(DataFile).filter_by(company_id=company.id).count() == 0

    def test_sync_idempotent_keeps_status(self, session, company, tmp_path):
        _touch(tmp_path, company.code, 2024, "BCQT", "Mau15_NVL.xlsx")
        sync_data_files(session, company, raw_root=tmp_path)
        row = session.query(DataFile).filter_by(company_id=company.id).first()
        row.parse_status = DataFileStatus.OK
        row.row_count = 42
        session.commit()

        sync_data_files(session, company, raw_root=tmp_path)  # re-sync
        row = session.query(DataFile).filter_by(company_id=company.id).first()
        assert row.parse_status == DataFileStatus.OK  # giữ trạng thái
        assert row.row_count == 42


class TestRecordParseResult:
    def test_record_marks_ok_with_counts(self, session, company, tmp_path):
        _touch(tmp_path, company.code, 2024, "BCQT", "Mau15_NVL.xlsx")
        _touch(tmp_path, company.code, 2024, "DINH_MUC", "BCDM_TT39.xls")
        sync_data_files(session, company, raw_root=tmp_path)

        stats = IngestStats(company_code=company.code, period_year=2024,
                            m15_rows=99, m16_rows=476)
        record_parse_result(session, company, 2024, stats)

        by_slot = {r.slot: r for r in session.query(DataFile).filter_by(company_id=company.id).all()}
        assert by_slot["m15"].parse_status == DataFileStatus.OK
        assert by_slot["m15"].row_count == 99
        assert by_slot["m16"].row_count == 476

    def test_record_writes_match_source_for_every_slot(self, session, company, tmp_path):
        """`match_source` phải xuống CỘT cho mọi slot, không chỉ Mẫu 15/15a (#84).

        Trước đây Mẫu 16 và tờ khai không sinh khoá đó ở provenance, nên cột rỗng và
        khối căn cứ đọc ở trang file không có gì để hiện.
        """
        from app.adapters._common import ParseProvenance
        from app.adapters.evidence import HEADER_MATCHED

        _touch(tmp_path, company.code, 2024, "DINH_MUC", "BCDM_TT39.xls")
        _touch(tmp_path, company.code, 2024, "HANG_CHI_TIET", "BCCT_NK.xlsx")
        sync_data_files(session, company, raw_root=tmp_path)

        stats = IngestStats(company_code=company.code, period_year=2024,
                            m16_rows=10, bcct_rows=20)
        stats.provenance = {
            "m16": ParseProvenance(
                detail={"form_signature": "sig-m16", "column_map": {"norm_qty": 7},
                        "template_id": None, "match_source": "keyword"},
                evidence={"norm_qty": HEADER_MATCHED},
            ),
            "bcct": ParseProvenance(
                detail={"form_signature": "sig-bcct", "column_map": {"quantity": 26},
                        "template_id": None, "match_source": "keyword"},
                evidence={"quantity": HEADER_MATCHED},
            ),
        }
        record_parse_result(session, company, 2024, stats)

        by_slot = {r.slot: r for r in session.query(DataFile).filter_by(company_id=company.id).all()}
        assert by_slot["m16"].match_source == "keyword"
        assert by_slot["bcct"].match_source == "keyword"

    def test_record_zero_rows_is_error(self, session, company, tmp_path):
        # WARNING không còn là status lifecycle (ADR #18): 0 dòng = chưa dùng được → error.
        _touch(tmp_path, company.code, 2024, "BCQT", "Mau15_NVL.xlsx")
        sync_data_files(session, company, raw_root=tmp_path)
        record_parse_result(session, company, 2024,
                            IngestStats(company_code=company.code, period_year=2024, m15_rows=0))
        row = session.query(DataFile).filter_by(company_id=company.id, slot="m15").first()
        assert row.parse_status == DataFileStatus.ERROR


def test_files_by_year_slot_groups(session, company, tmp_path):
    _touch(tmp_path, company.code, 2024, "BCQT", "Mau15_NVL.xlsx")
    _touch(tmp_path, company.code, 2024, "HANG_CHI_TIET", "BCCT_NK.xlsx")
    _touch(tmp_path, company.code, 2024, "HANG_CHI_TIET", "BCCT_XK.xlsx")
    sync_data_files(session, company, raw_root=tmp_path)

    grouped = files_by_year_slot(session, company)
    assert set(grouped[2024].keys()) == {"m15", "bcct"}
    assert len(grouped[2024]["bcct"]) == 2  # nhiều file BCCT


# Nhãn "Đã nạp · X/4 loại" đã bị bỏ ở #86: đủ 4 loại không có nghĩa là kiểm tra chạy
# được, nên thước đo là "đủ dữ liệu cho N/M kiểm tra" (ADR #24 mục 1). Bất biến thay
# thế nằm ở `tests/test_data_screen.py` và `tests/test_readiness.py`.


def test_one_workbook_registers_under_every_slot_it_serves(tmp_path):
    """Một DN gộp Mẫu 15/15a/16 vào MỘT workbook, tên file không nói được biểu nào.

    Registry phải khớp `discover`; lệch thì trang tài liệu báo "chưa có file" trong
    khi dữ liệu của chính file đó đã nạp.
    """
    import openpyxl

    from app.pipeline.data_files import _iter_company_files

    def sheet(ws, header, code_prefix):
        for _ in range(8):
            ws.append([None] * len(header))
        ws.append(header)
        for i in range(2):
            ws.append([i + 1, f"{code_prefix}{i}", "Tên", "KG", 10, 100, 0, 0, 80, 0, 30])

    d = tmp_path / "DN_G" / "2025" / "BCQT"
    d.mkdir(parents=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet(wb.create_sheet("15-BCQT-NVL"),
          ["STT", "Mã NVL", "Tên NVL", "Đơn vị tính", "Tồn đầu kỳ", "Nhập trong kỳ",
           "Tái xuất", "Chuyển MĐSD", "Xuất sản xuất", "Xuất khác", "Tồn cuối kỳ"], "MAT")
    sheet(wb.create_sheet("TP-15a"),
          ["STT", "Mã SP", "Tên SP", "Đơn vị tính", "Tồn đầu kỳ", "Nhập kho trong kỳ",
           "Chuyển MĐSD", "Xuất khẩu", "Xuất khác", "Tồn cuối kỳ"], "SP")
    wb.save(d / "Báo cáo quyết toán 2025.xlsx")

    slots = {slot for _y, slot, _p in _iter_company_files(tmp_path, "DN_G")}
    # Tên file khớp không rule nào -> trước đây bị bỏ hẳn, registry rỗng.
    assert "m15" in slots
    assert "m15a" in slots
