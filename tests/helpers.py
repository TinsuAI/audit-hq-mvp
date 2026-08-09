"""Tiện ích dùng chung cho test HTTP.

`TestClient(app)` không vào context manager nên lifespan không chạy → không có
worker thread. Từ khi nạp dữ liệu chuyển sang hàng đợi (job `ingest`), test phải
tự chạy job đã xếp thì mới thấy trạng thái sau khi nạp.
"""

from __future__ import annotations

import io
import re
from functools import lru_cache
from pathlib import Path

import app.database as dbmod

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

M15_HEADER = [
    "STT", "Mã NVL", "Tên NVL", "Đơn vị tính", "Tồn đầu kỳ", "Nhập trong kỳ",
    "Tái xuất", "Chuyển mục đích sử dụng", "Xuất sản xuất", "Xuất khác", "Tồn cuối kỳ",
]


def m15_xlsx_bytes(header: list[str] = M15_HEADER, rows: int = 3) -> bytes:
    """Workbook Mẫu 15 tối thiểu: 8 dòng đầu trống, dòng tiêu đề, rồi `rows` dòng NVL."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "BCQT_NVL"
    for _ in range(8):
        ws.append([None] * len(header))
    ws.append(header)
    for i in range(rows):
        ws.append([i + 1, f"MAT{i}", "Tên", "KG", 10, 100, 0, 0, 80, 0, 30])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def unreadable_m15_xlsx_bytes(header: list[str] = M15_HEADER) -> bytes:
    """Workbook mà `select_sheet` KHÔNG nhận ra trang nào đúng biểu.

    Dựng lại tình huống file cán bộ tự gộp: chèn một cột đầu (nhãn nguồn) và bỏ khối
    tiêu đề → mọi cột lệch một ô, mọi trang chấm 0 điểm, file bị từ chối khi nạp.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Bìa"
    ws.append(["Trang bìa"])
    ws2 = wb.create_sheet("Dữ liệu")
    ws2.append(["Nguồn", *header])
    for i in range(3):
        ws2.append(["F1", i + 1, f"MAT{i}", "Tên", "KG", 10, 100, 0, 0, 80, 0, 30])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload_and_ingest(
    client, code: str, name: str, content: bytes, year: int = 2024, mime: str = XLSX_MIME
):
    """Đường nạp qua web sau #88: thả file vào ô thả, RỒI bấm nút nạp.

    Ô thả không xếp việc nạp — cán bộ còn phải sửa loại và gán sổ trước. Trả phản hồi
    của bước nạp (303 về trang công việc), người gọi tự `drain_jobs()`.
    """
    r = client.post(
        f"/companies/{code}/upload",
        data={"year": str(year)},
        files=[("files", (name, content, mime))],
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text
    return client.post(
        f"/companies/{code}/documents/ingest",
        data={"year": str(year)},
        follow_redirects=False,
    )


def drain_jobs(max_rounds: int = 5) -> int:
    """Chạy hết job đang `queued`, kể cả job do job khác xếp thêm. Trả số job đã chạy.

    `ingest` xếp tiếp `run_checks` khi payload yêu cầu, nên phải lặp chứ không
    quét một lượt.
    """
    from app.jobs import HANDLERS, register_handler, run_job
    from app.jobs.handlers import ingest_handler, run_batch_handler, run_checks_handler
    from app.models.job import Job, JobKind, JobStatus

    for kind, fn in (
        (JobKind.RUN_CHECKS, run_checks_handler),
        (JobKind.BATCH_RUN, run_batch_handler),
        (JobKind.INGEST, ingest_handler),
    ):
        if kind.value not in HANDLERS:
            register_handler(kind, fn)

    ran = 0
    for _ in range(max_rounds):
        with dbmod.SessionLocal() as db:
            jobs = (
                db.query(Job)
                .filter_by(status=JobStatus.QUEUED.value)
                .order_by(Job.id)
                .all()
            )
            if not jobs:
                return ran
            for job in jobs:
                run_job(db, job)
                ran += 1
    return ran


def last_job_result(kind: str | None = None) -> dict | None:
    """`result` của job mới nhất (lọc theo `kind` nếu cần) — để test đọc kết luận."""
    from app.models.job import Job

    with dbmod.SessionLocal() as db:
        q = db.query(Job)
        if kind:
            q = q.filter_by(kind=kind)
        job = q.order_by(Job.id.desc()).first()
        return dict(job.result) if job and job.result else None


# ─────────────────────── đọc mã nguồn lưới ───────────────────────

GRID_JS = Path(__file__).resolve().parents[1] / "app" / "static" / "cell-grid.js"


@lru_cache(maxsize=1)
def grid_source() -> str:
    return GRID_JS.read_text(encoding="utf-8")


@lru_cache(maxsize=1)
def grid_functions() -> dict[str, str]:
    """Thân từng hàm mức trên cùng của `cell-grid.js`, cắt theo cột thụt đầu dòng.

    Cắt thô có chủ ý: bộ test không nhúng trình phân tích JS, và một regex chặt hơn thì
    mọi lần sửa `cell-grid.js` sau này đỏ vì lý do không liên quan.

    Dùng chung vì nhiều vé cần cùng phép đọc này: mã chỉ chạy ở trình duyệt đứt im lặng,
    nên cách bắt được nó là lấy bộ chọn / id TỪ JS rồi đối chiếu với trang đã dựng.
    """
    src = grid_source()
    starts = [(m.group(1), m.start()) for m in re.finditer(r"^  function (\w+)\(", src, re.M)]
    return {
        name: src[pos : (starts[i + 1][1] if i + 1 < len(starts) else len(src))]
        for i, (name, pos) in enumerate(starts)
    }


def grid_body(name: str) -> str:
    fns = grid_functions()
    assert name in fns, f"`cell-grid.js` không còn hàm `{name}`"
    return fns[name]
